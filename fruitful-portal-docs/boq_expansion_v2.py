# boq_expansion_v2.py
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- theme (match existing) ---
FONT, HEADER_FILL, HEADER_TXT, BAND, BORDER = 'Arial', '0FB5A6', 'FFFFFF', 'F2F2F2', 'BFBFBF'

new_data = [
    # --- PPE & FOOD SAFETY (The "McDonald's Standard" for zero-contamination) ---
    ["PPE & Safety", "Food-Grade Latex-Free Gloves", "Powder-Free, Textured Grip", "Daily", "Stock S/M/L; High-sensitivity for artisan work"],
    ["PPE & Safety", "Beard Covers & Hair Nets", "Breathable, Branded Band", "Daily", "Mandatory for all food prep areas"],
    ["PPE & Safety", "Disposable Poly-Aprons", "Clear, 20-micron", "Daily", "Overlay for high-mess pancake prep"],
    ["PPE & Safety", "First Aid Kit - Refill Pack", "SABS/Dept of Labour", "Monthly", "Antiseptics, burn-gel, blue detectable plasters"],

    # --- CONDIMENT BAR & GUEST NEEDS (The "Starbucks Standard") ---
    ["Condiment Bar", "Individually Wrapped Sugar", "Branded White/Brown/Sweetener", "Daily", "MOQ 100k; Moisture-proof paper"],
    ["Condiment Bar", "Artisan Spice Shakers", "Cinnamon/Cocoa/Nutmeg", "Weekly", "Branded glass with stainless swivel tops"],
    ["Condiment Bar", "Disposable Stirrers & Straws", "FSC Wood / Paper-wrapped", "Daily", "Branded sleeves; wide-bore for smoothies"],
    ["Condiment Bar", "Guest Napkin Dispensers", "Countertop / Branded", "Permanent", "One-at-a-time pull to reduce waste"],
    ["Condiment Bar", "Wet-Wipes (Branded)", "Lemon/Alcohol-free", "Daily", "For pancake 'finger-food' cleanup"],

    # --- BACK-OF-HOUSE (BOH) UTILITIES (The "Invisible" needs) ---
    ["BOH Utilities", "Heavy-Duty Pancake Tongs", "Silicone-Tipped / Branded", "Permanent", "Heat resistant; long-handle for kiosk safety"],
    ["BOH Utilities", "Batter Mixing Buckets", "NSF Graduated (10L/20L)", "Permanent", "Food-grade HDPE with lids"],
    ["BOH Utilities", "Whisks & Scrapers", "Commercial Grade / High-torque", "Permanent", "Balloon whisks and stainless bench scrapers"],
    ["BOH Utilities", "Kitchen Timers (Magnetic)", "Rugged Digital", "Permanent", "Loud-alert; waterproof membrane"],
    ["BOH Utilities", "Sharpie/Marker Set (Food Safe)", "Industrial Black/Red", "Monthly", "For marking 'Use By' dates on canisters"],

    # --- SANITATION & WASTE (365/24/7 Cycle) ---
    ["Sanitation", "Degreaser - Surface Spray", "SABS Approved / Food Safe", "Weekly", "Trigger bottles + 5L bulk refill"],
    ["Sanitation", "Dish Soap - Antibacterial", "High-Concentrate", "Weekly", "For smallwares soaking"],
    ["Sanitation", "Trash Bag Liners (Heavy Duty)", "Recycled / 30L & 50L", "Daily", "High-puncture resistance for kiosk bins"],
    ["Sanitation", "Sanitizer Buckets (Red/Green)", "Color-coded 5L", "Permanent", "Standardized for 'Quat' sanitizer solution"],

    # --- KIOSK OPS (The "Closing" Pack) ---
    ["Kiosk Ops", "Cash Float Bags", "Tamper-Evident / Branded", "Daily", "For bank-mandatory deposit protocols"],
    ["Kiosk Ops", "Closing Checklists (Laminated)", "Fruitful SOP v1.2", "Permanent", "Use with dry-erase markers"],
    ["Kiosk Ops", "Lock-out/Tag-out (LOTO) Kit", "OSHA Standard", "Permanent", "Safety for machine maintenance"],
    ["Kiosk Ops", "Spare Thermal Paper (POS)", "80mm / Long-roll", "Monthly", "BPA-free; high-contrast print"]
]

wb = load_workbook('Fruitful_Portal_Master_BOQ_Enhanced.xlsx')
ws = wb.active
start_row = ws.max_row + 1

thin = Side(style='thin', color=BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
band_fill = PatternFill('solid', start_color=BAND)

for i, row_data in enumerate(new_data, start=start_row):
    for j, val in enumerate(row_data, 1):
        c = ws.cell(i, j, val); c.border = border
        if (i % 2) == 1: c.fill = band_fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save('Fruitful_Portal_Master_BOQ_Exhaustive.xlsx')
print("DONE")
