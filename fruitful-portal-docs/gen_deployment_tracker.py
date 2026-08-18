# gen_deployment_tracker.py — merge mall register + tier economics into full deployment tracker
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont = Font(name=FONT, size=11, bold=True, color='FFFFFF')
hfill = PatternFill('solid', start_color='0FB5A6')
bfill = PatternFill('solid', start_color='F2F2F2')
gold = PatternFill('solid', start_color='FFF2CC')
green = PatternFill('solid', start_color='E2F0D9')

def sheet(ws, headers, widths):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h); c.font = hfont; c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'C2'

def put(ws, r, vals, fill=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = border
        if r % 2 == 1 and fill is None: c.fill = bfill
        if fill: c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

def money(v): return f"R{v:,.0f}"

# ---------- tier economics (from mall deployment model) ----------
ECON = {
    "Premium":       dict(rev=400000, rent=20000, lab=30000, jura=19482, opex=6000, deploy=85000, deposit=40000, rates=2500),
    "Regional":      dict(rev=280000, rent=12000, lab=25000, jura=15784, opex=5000, deploy=70000, deposit=24000, rates=1750),
    "Neighbourhood": dict(rev=180000, rent=6000,  lab=22000, jura=10086, opex=4000, deploy=55000, deposit=12000, rates=1150),
}
LANDED = 123085

def econ(tier):
    e = ECON[tier]
    cogs = e["rev"] * 0.45
    prof = e["rev"] - cogs - e["lab"] - e["rent"] - e["jura"] - e["opex"] - e["rates"]
    capex = LANDED + e["deploy"] + e["deposit"]
    return cogs, prof, capex

# ---------- read register ----------
wb_src = load_workbook('FRUITFUL_MALL_REGISTER.xlsx')
ws_src = wb_src['Mall Register (ranked)']
malls = []
for row in ws_src.iter_rows(min_row=2, values_only=True):
    malls.append({"rank": row[0], "name": row[1], "addr": row[2], "prov": row[3], "city": row[4],
                  "km": row[5], "tier": row[6], "email": row[7], "wave": row[8]})

# ---------- tracker workbook ----------
wb = Workbook()
ws = wb.active; ws.title = 'Deployment Tracker'
ws.sheet_properties.tabColor = '0FB5A6'
headers = ["Rank", "Mall / Centre", "Tier", "City", "km", "Wave", "Est rent/mo", "Est profit/mo",
           "Capex (landed+deploy+dep)", "Payback (mo)", "Contact email",
           "Lease status", "Rent agreed (R/mo)", "NUCO deploy date", "Go-live date", "Notes"]
sheet(ws, headers, [6, 34, 13, 14, 7, 11, 12, 13, 15, 11, 30, 16, 14, 13, 12, 24])
r = 2
for m in malls:
    tier = m["tier"] if m["tier"] in ECON else "Neighbourhood"
    cogs, prof, capex = econ(tier)
    payback = capex / prof
    put(ws, r, [m["rank"], m["name"], tier, m["city"] or "", m["km"], m["wave"] or "",
                money(ECON[tier]["rent"]), money(prof), money(capex), f"{payback:.1f}",
                m["email"], "Not approached", "", "", "", ""])
    r += 1

# summary sheet
ws2 = wb.create_sheet('Summary')
sheet(ws2, ["Metric", "Value"], [44, 26])
from collections import Counter
tc = Counter(m["tier"] for m in malls)
waves = Counter(m["wave"] or "Unassigned" for m in malls)
tot_prof = sum(econ(m["tier"] if m["tier"] in ECON else "Neighbourhood")[1] for m in malls)
tot_capex = sum(econ(m["tier"] if m["tier"] in ECON else "Neighbourhood")[2] for m in malls)
summ = [
    ("Total centres", len(malls)),
    ("Premium", tc.get("Premium", 0)), ("Regional", tc.get("Regional", 0)), ("Neighbourhood", tc.get("Neighbourhood", 0)),
    ("Wave 1 (PTA/JHB)", waves.get("Wave 1 (PTA/JHB)", 0) + waves.get("Wave 1b", 0)),
    ("Wave 2", waves.get("Wave 2", 0)), ("Wave 3+", waves.get("Wave 3+", 0)),
    ("Combined est. profit/mo (all 117 live)", money(tot_prof)),
    ("Combined est. capex (all 117)", money(tot_capex)),
    ("Payback weighted (mo)", f"{tot_capex/tot_prof:.1f}"),
]
for i, (k, v) in enumerate(summ, 2):
    put(ws2, i, [k, v], fill=gold if k.startswith("Combined") else None)

wb.save('FRUITFUL_DEPLOYMENT_TRACKER.xlsx')
print(f"tracker rows: {len(malls)}")
