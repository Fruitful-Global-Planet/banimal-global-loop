# gen_business_model.py — FRUITFUL GLOBAL BUSINESS MODEL (businessman rebuild)
# Fixes the withdrawn "fully-loaded per-unit" presentation (which put ~R10.13 of
# store-level fixed cost onto a R1.90 water pancake -> -485% "margin").
#
# How real franchise systems price (McDonald's / Starbucks licensed):
#   - Per-unit (menu engineering): price - direct cost = CONTRIBUTION. Stop there.
#   - Royalty + marketing = % of STORE REVENUE, charged once per month per store
#     (McDonald's royalty 4-5% + ad fund ~2-4% of gross sales; Starbucks licensed
#     ~6-8% of sales). Same % at location #1 and location #100,000. Never per-item.
#   - Fixed costs (labour, rent, machine) sit in the store P&L, recovered from the
#     aggregate contribution of all units sold, not added into each unit's cost.
#
# Structure: Unit contribution -> Store P&L -> HQ fee pool -> Global price map.
# FX: live rates via open.er-api.com (free), fetched at run time.
import json
import urllib.request

FX = 16.185472  # USD->ZAR, Sun 16 Aug 2026 00:02 UTC (live fetch at build)
try:
    with urllib.request.urlopen('https://open.er-api.com/v6/latest/USD', timeout=15) as r:
        d = json.load(r)
    USD_RATES = d['rates']
    FX = USD_RATES['ZAR']
    FX_NOTE = f"FX: live from exchangerate-api.com, {d['time_last_update_utc']}; 1 USD = R{FX:.6f}"
except Exception:
    USD_RATES = None
    FX_NOTE = f"FX: 1 USD = R{FX:.6f} (16 Aug 2026; live fetch unavailable at build)"

def fx_to_zar(ccy):
    # foreign units per 1 ZAR (same convention as FRUITFUL_GLOBAL_PORTAL_MATRIX: USD 0.06178)
    if ccy == 'ZAR':
        return 1.0
    return USD_RATES[ccy] / FX if USD_RATES else None

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont = Font(name=FONT, size=11, bold=True, color='FFFFFF')
hfill = PatternFill('solid', start_color='0FB5A6')
bfill = PatternFill('solid', start_color='F2F2F2')
gold = PatternFill('solid', start_color='FFF2CC')
grey = Font(name=FONT, size=10, italic=True, color='666666')

def usd(v):
    if isinstance(v, str) or v is None:
        return v
    return f"${v / FX:,.2f}"

def zar(v):
    if isinstance(v, str) or v is None:
        return v
    return f"R{v:,.2f}"

def sheet(ws, headers, widths):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h); c.font = hfont; c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'

def put(ws, r, vals, fill=None, bold=False):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = border
        if r % 2 == 1 and fill is None: c.fill = bfill
        if fill: c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if bold:
            c.font = Font(name=FONT, size=11, bold=True)

def note(ws, text):
    c = ws.cell(ws.max_row + 1, 1, text)
    c.font = grey
    c.alignment = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ================= Sheet 1: Unit Contribution (menu engineering) =================
ws = wb.active; ws.title = 'Unit Contribution'
ws.sheet_properties.tabColor = '0FB5A6'
sheet(ws, ["P/N", "SKU", "Spec (one version only)", "Price ZAR", "Direct ZAR",
           "Contrib ZAR", "Contrib %", "Price USD", "Direct USD", "Contrib USD"],
      [10, 26, 34, 11, 11, 12, 10, 11, 11, 12])
# price_zar, direct_zar
LINES = [
    ("FFP-01", "The Coffee with Milk", "350ml, 18g beans, 150ml milk", 40.00, 14.66),
    ("FFP-02", "The Coffee with Froth", "350ml, +foam (no extra cost)", 40.00, 14.66),
    ("FFP-03", "The Coffee with Chocolate", "350ml, +8g cocoa", 40.00, 17.06),
    ("FFP-04", "Green Energy smoothie", "400ml one size", 60.00, 23.40),
    ("FFP-05", "Mango Lassi smoothie", "400ml one size", 60.00, 23.40),
    ("FFP-06", "Power Protein smoothie", "400ml one size", 60.00, 23.40),
    ("FFP-07", "Pannekoek milk 100% (boxed)", "1pc, cinnamon sugar + lemon", 12.50, 3.71),
    ("FFP-07x12", "Pannekoek milk 12-pack", "12 units, sealed box", 150.00, 44.52),
    ("FFP-09", "Pannekoek water 50% (no box)", "half size, impulse", 1.90, 0.75),
    ("FFP-09x12", "Pannekoek water 12-pack", "12 units, no boxes", 22.50, 9.00),
]
r = 2
for pn, sku, spec, p, c in LINES:
    contrib = p - c
    put(ws, r, [pn, sku, spec, zar(p), zar(c), zar(contrib), f"{contrib/p*100:.1f}%",
                usd(p), usd(c), usd(contrib)])
    r += 1
put(ws, r, ["", "Coffee blended (FFP-01..03, equal mix)", "avg direct R15.46 -> R24.54 contrib",
            "", "", "", "61.4%", "", "", ""], fill=gold)
note(ws, "Menu-engineering rule (McDonald's / Starbucks): the per-unit view STOPS at contribution. "
         "No labour, rent, royalty or marketing is ever added into unit cost. Those live in the "
         "store P&L as % of revenue and fixed costs. The old 'fully loaded per-unit' column is "
         "withdrawn: it added R10.13/unit of store-level cost onto a R1.90 item (-485% 'margin') - "
         "an artifact of the allocation method, not a real cost. Water pancake is an impulse SKU "
         "whose job is contribution at the margin; fixed costs are covered by total volume.")

# ================= Sheet 2: Store P&L per portal (monthly) =================
ws2 = wb.create_sheet('Store P&L per portal')
ws2.sheet_properties.tabColor = '0FB5A6'
sheet(ws2, ["Line", "Coffee (1,500/mo)", "Smoothies (750/mo)", "Pancakes (600/mo)", "TOTAL ZAR", "TOTAL USD"],
      [30, 16, 16, 16, 14, 14])
COF_Q, SM_Q, PAN_Q = 1500, 750, 600
COF_P, SM_P, PAN_P = 40.00, 60.00, 12.50          # final menu prices
COF_C, SM_C, PAN_C = 15.46, 23.40, 3.71            # blended coffee avg; smoothie; milk pancake
rev_c, rev_s, rev_p = COF_Q*COF_P, SM_Q*SM_P, PAN_Q*PAN_P
cog_c, cog_s, cog_p = COF_Q*COF_C, SM_Q*SM_C, PAN_Q*PAN_C
rev = rev_c + rev_s + rev_p
cogs = cog_c + cog_s + cog_p
gross = rev - cogs
pay = rev * 0.03
labour, fixed = 16800.00, 12082.00
roy, mkt = rev * 0.06, rev * 0.04
ebitda = gross - pay - labour - fixed - roy - mkt
dep = 4167.00
tax = (ebitda - dep) * 0.27
net_store = ebitda - dep - tax

def row2(name, z, fill=None, bold=False, money=True):
    vals = [name] + [zar(v) if money and isinstance(v, float) else v for v in z[:3]] + \
           ([zar(z[3]), usd(z[3])] if money else [z[3], z[3]])
    put(ws2, ws2.max_row + 1, vals, fill=fill, bold=bold)

row2("Units", [COF_Q, SM_Q, PAN_Q, COF_Q+SM_Q+PAN_Q], money=False)
row2("Revenue", [rev_c, rev_s, rev_p, rev])
row2("Direct COGS (grains above)", [cog_c, cog_s, cog_p, cogs])
row2("Gross profit", [rev_c-cog_c, rev_s-cog_s, rev_p-cog_p, gross], fill=gold, bold=True)
row2("Gross margin %", [f"{((rev_c-cog_c)/rev_c*100):.1f}%", f"{((rev_s-cog_s)/rev_s*100):.1f}%",
                        f"{((rev_p-cog_p)/rev_p*100):.1f}%", f"{gross/rev*100:.1f}%"], money=False)
row2("Payment fees (3% of revenue)", ["", "", "", pay])
row2("Labour - 2 staff (manager = HQ)", ["", "", "", labour])
row2("Mall rent + Jura + opex", ["", "", "", fixed])
row2("Royalty (6% of revenue, to HQ)", ["", "", "", roy])
row2("Marketing fund (4% of revenue, to HQ)", ["", "", "", mkt])
row2("EBITDA", ["", "", "", ebitda], fill=gold, bold=True)
row2("Depreciation (R4,167) + tax 27%", ["", "", "", -(dep + tax)])
row2("NET PROFIT / month", ["", "", "", net_store], fill=gold, bold=True)

be_contrib = (gross - pay - roy - mkt) / (COF_Q + SM_Q + PAN_Q)   # contribution after variable fees
be_fixed = labour + fixed
be_units = be_fixed / be_contrib
be_rev = be_units * (rev / (COF_Q + SM_Q + PAN_Q))
put(ws2, ws2.max_row + 2, ["BREAK-EVEN (cash, after variable fees)", "", "",
                           f"{be_units:,.0f} units/mo  ~{be_units/30:,.0f}/day", zar(be_rev), usd(be_rev)],
    fill=gold, bold=True)
note(ws2, f"Net R{net_store:,.2f}/mo (${net_store/FX:,.2f}) at FINAL prices (R40/R60/R12.50). The earlier R21,554 "
          f"baseline used the pre-final R35 pancake price; final pancake pricing is R12.50 (milk) / "
          f"R1.90 (water), with 12-packs (R150/R22.50) as the DC/retail unit on top.")
note(ws2, "Break-even: fixed cash R28,882/mo (labour + mall/Jura/opex) divided by contribution after "
          "variable fees (payment 3% + royalty 6% + marketing 4% all scale with revenue). % fees "
          "never change the price of a single item - they are charged once on the till total.")
note(ws2, FX_NOTE)

# ================= Sheet 3: HQ Fee Pool (franchise layer) =================
ws3 = wb.create_sheet('HQ Fee Pool')
ws3.sheet_properties.tabColor = '0FB5A6'
sheet(ws3, ["Line", "Per portal / month ZAR", "Per portal / month USD"], [44, 22, 22])
fees = [
    ("Store revenue (baseline)", rev, "coffee 1,500 + smoothies 750 + pancakes 600"),
    ("Royalty 6% of revenue", roy, "franchisor % fee - same at location 1 and 100,000"),
    ("Marketing fund 4% of revenue", mkt, "national campaigns, funded by system"),
    ("Fee pool to HQ", roy + mkt, "10% of store revenue"),
    ("HQ fixed: manager", 10000.00, "paid from the pool, NOT from portal labour"),
    ("HQ net (1 portal)", roy + mkt - 10000.00, "funds HQ expansion, support, R&D"),
]
r = 2
for name, z, cmt in fees:
    put(ws3, r, [name, zar(z), usd(z)])
    put(ws3, r + 0, None) if False else None
    r += 1
put(ws3, r, ["", "", ""])
put(ws3, r, ["PORTALS IN SYSTEM", "Pool to HQ ZAR", "Pool to HQ USD"], bold=True)
r += 1
for n in [1, 5, 10, 50, 100]:
    pool = (roy + mkt) * n
    net_hq = pool - 10000
    put(ws3, r, [f"{n} portal{'s' if n > 1 else ''}", zar(net_hq), usd(net_hq)])
    r += 1
note(ws3, "Franchisor model: % of store sales, applied once per store per month - never a per-item "
          "surcharge. Reference points: McDonald's royalty 4-5% of gross sales + advertising fund "
          "~2-4% (~9% total); Starbucks licensed stores ~6-8% of sales. Fruitful at 10% (6+4) sits "
          "in the same band and funds the manager + national marketing from the system, not from "
          "any single portal's unit costs.")
note(ws3, FX_NOTE)

# ================= Sheet 4: Global Price Map (9 currencies, final menu) =================
ws4 = wb.create_sheet('Global Price Map')
ws4.sheet_properties.tabColor = '0FB5A6'
MARKETS = [  # same 9 currencies as FRUITFUL_GLOBAL_PORTAL_MATRIX
    ("South Africa", "ZAR", "R", 2), ("USA / Zimbabwe", "USD", "$", 2),
    ("Japan", "JPY", "¥", 0), ("Botswana", "BWP", "P", 2),
    ("Mexico", "MXN", "$", 2), ("Iceland", "ISK", "kr", 0),
    ("Egypt", "EGP", "EGP ", 2), ("Thailand", "THB", "฿", 0),
    ("Canada", "CAD", "$", 2),
]
sheet(ws4, ["Market", "Ccy", "FX (ZAR->local)", "Coffee R40", "Smoothie R60",
            "Pancake milk R12.50", "Pancake water R1.90", "Milk 12pk R150", "Water 12pk R22.50"],
      [22, 8, 14, 13, 13, 15, 15, 13, 15])
BASE = {"Coffee": 40.00, "Smoothie": 60.00, "Milk single": 12.50, "Water single": 1.90,
        "Milk 12pk": 150.00, "Water 12pk": 22.50}
r = 2
for market, ccy, sym, dec in MARKETS:
    f = fx_to_zar(ccy)
    if f is None:
        continue
    def px(z):
        v = z * f
        return f"{sym}{v:,.{dec}f}" if dec else f"{sym}{v:,.0f}"
    put(ws4, r, [market, ccy, f"{f:.4f}"] + [px(BASE[k]) for k in BASE])
    r += 1
note(ws4, "Final radical-menu prices converted at live FX (exchangerate-api, 16 Aug 2026). "
          "One price per category, no sizes - same structure in every market. Prices shown are "
          "FX conversions of the SA prices; local market positioning (rounding, tax, rent index) "
          "is a pricing decision, not a cost input.")
note(ws4, FX_NOTE)

wb.save('FRUITFUL_BUSINESS_MODEL_GLOBAL.xlsx')
print("DONE - FRUITFUL_BUSINESS_MODEL_GLOBAL.xlsx")
print(FX_NOTE)
print(f"\nStore P&L: rev R{rev:,.2f} | cogs R{cogs:,.2f} | gross R{gross:,.2f} ({gross/rev*100:.1f}%)")
print(f"  fees: pay {pay:,.2f} | labour {labour:,.2f} | fixed {fixed:,.2f} | roy {rev*0.06:,.2f} | mkt {rev*0.04:,.2f}")
print(f"  EBITDA R{ebitda:,.2f} | dep+tax R{dep+tax:,.2f} | NET R{net_store:,.2f} (${net_store/FX:,.2f})")
print(f"  Break-even: {be_units:,.0f} units/mo (~{be_units/30:,.0f}/day) = R{be_rev:,.0f}/mo (${be_rev/FX:,.0f})")
print(f"HQ pool/portal: R{roy+mkt:,.2f} (${(roy+mkt)/FX:,.2f}); manager R10,000; net at 5 portals R{(roy+mkt)*5-10000:,.2f}")
