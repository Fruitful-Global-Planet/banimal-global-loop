from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- theme ---
FONT, HEADER_FILL, HEADER_TXT, BAND, BORDER = 'Arial', '0FB5A6', 'FFFFFF', 'F2F2F2', 'BFBFBF'

data = [
    ["Kiosk Hardware", "1 pc Sample (shipping)", "Guangzhou Welink Smart Tech", "Welink WL20260325 PI", "Track arrival; verify 32\"/43\" screen res"],
    ["Kiosk Hardware", "Fleet Target: 300 units", "TBD (Welink vs Alt)", "China Order Analysis", "Find 2 alternative solar-kiosk suppliers for price/spec comparison"],
    ["Consumables", "Artisan Coffee Beans", "Fantastico SA (Auth)", "Coffee Trading Pack", "Confirm SCA Score 80+ / Moisture 9-12% on first lot"],
    ["Consumables", "Branded Cups (8/10/12oz)", "TBD", "Fruitful CI Guide v04", "Find factory for branded double-wall cups; MOQ 50k"],
    ["Staff Apparel", "Uniform Line (Caps to Shoes)", "TBD", "Staff & Ops Spec (2026-2030)", "Source suppliers for Vellies (embossed), Bomber Jackets, and Polo Shirts"],
    ["Staff Apparel", "Safety Shoes (Steel-toe)", "TBD", "Audit Registry", "Source branded SABS-approved safety footwear"],
    ["IT / Tech", "Rugged 10\" Staff Tablets", "TBD", "Staff & Ops Spec", "Find industrial Android tablets with 12hr+ battery"],
    ["IT / Tech", "Digital e-Ink Nametags", "TBD", "base44 IAM Spec", "Source e-Ink display tags with NFC/RFID and API sync"],
    ["IT / Tech", "Digital Writing Pads", "TBD", "Staff & Ops Spec", "Source e-Ink tablets for clean-syncing logbooks"],
    ["IT / Tech", "Industrial 5G Routers", "TBD", "Audit Registry", "Source high-temp rugged 5G/LTE routers with external antenna"],
    ["Facilities", "Jura Cleaning Tablets", "Jura / Generic", "Audit Registry", "Compare OEM vs verified generic cost for bulk (300 portals)"],
    ["Facilities", "Biometric Attendance Sys", "TBD", "Audit Registry", "Source cloud-synced facial/fingerprint hardware"]
]

cols = ["Category", "Item/Target", "Supplier Status", "Authoritative Spec (HAVE)", "Agent Action (NEED)"]

wb = Workbook()
ws = wb.active; ws.title = 'Activation Checklist'
ws.sheet_properties.tabColor = HEADER_FILL

thin = Side(style='thin', color=BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name=FONT, size=11, bold=True, color=HEADER_TXT)
header_fill = PatternFill('solid', start_color=HEADER_FILL)
band_fill   = PatternFill('solid', start_color=BAND)

for j, name in enumerate(cols, 1):
    c = ws.cell(1, j, name); c.font = header_font; c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border

for i, row_data in enumerate(data, start=2):
    for j, val in enumerate(row_data, 1):
        c = ws.cell(i, j, val); c.border = border
        if (i % 2) == 1: c.fill = band_fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

for j, name in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(j)].width = 30

ws.freeze_panes = 'A2'
wb.save('Fruitful_Sourcing_Activation_Checklist.xlsx')
print("DONE")
