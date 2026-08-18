# boq_final.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT, HEADER_FILL, HEADER_TXT, BAND, BORDER = 'Arial', '0FB5A6', 'FFFFFF', 'F2F2F2', 'BFBFBF'

data = [
    # ============ A. SPILL RESPONSE & SURFACE CARE (coffee spills!) ============
    ["Spill Response", "Coffee Spill Wipes — Instant", "Absorbent cellulose, 100ct tub", "Daily", "Wipe spills on contact; blue-tint detectable"],
    ["Spill Response", "Bar Towels (Microfibre)", "Branded, 40×60cm, 20-pack", "Weekly", "Colour-coded: green = counter, blue = glass"],
    ["Spill Response", "Wet Floor Sign / Cones", "Folding A-frame", "Permanent", "SA H&S requirement — immediate spill safety"],
    ["Spill Response", "Spot Cleaner Spray (Stain)", "Multi-surface, food-safe", "Weekly", "For kiosk body fibreglass & counter marks"],
    ["Spill Response", "Absorbent Floor Granules", "Oil/grease absorb, 10kg bag", "Monthly", "For dropped milk/syrup slicks"],

    # ============ B. HEAD OFFICE & COMMUNICATIONS ============
    ["Comms", "Head Office Hotline — VoIP Desk Phone", "VOIP handset + headset", "Permanent", "Dedicated line to Silver Lakes HQ; speed-dial 1"],
    ["Comms", "WhatsApp Business Line (Staff)", "Dedicated e-SIM tablet/phone", "Permanent", "+27 72 414 0117 escalation line; shift-report group"],
    ["Comms", "Two-Way Radios (Kiosk Pair)", "Long-range UHF, rechargeable", "Permanent", "Kiosk counter ↔ BOH/mobile operator"],
    ["Comms", "Incident Report Pads", "Carbonless, 2-part, branded", "Quarterly", "Paper backup for the digital incident log"],
    ["Comms", "Emergency Contact Card (Laminated)", "HQ, fire, security, medical", "Permanent", "By the phone — staff first-day pack item"],

    # ============ C. CASH & HARD CURRENCY ============
    ["Cash Handling", "Cash Register Drawer (Lockable)", "Metal, 4-coin/6-bill, teal", "Permanent", "Bank-mandatory; mounts in stainless cabinet"],
    ["Cash Handling", "Cash Box (Drop-Safe, Tamper-Evident)", "Steel, key + combo, bolted", "Permanent", "Time-delay drop safe for mid-shift deposits"],
    ["Cash Handling", "Cash Float Bags (Branded)", "Tamper-evident zip, numbered", "Daily", "Bank-mandatory deposit protocol (per audit registry)"],
    ["Cash Handling", "Counterfeit-Detection Pens", "UV + chemical dual", "Monthly", "R50/R100/R200 note checks at shift end"],
    ["Cash Handling", "Coin Wrappers & Bank Straps", "ZAR denominations, 100pk", "Monthly", "Bank-ready cash packing"],
    ["Cash Handling", "Cash Register Receipts (Thermal 80mm)", "BPA-free, long-roll", "Monthly", "POS spare rolls — do not run dry"],
    ["Cash Handling", "Cash Pickup Service (SBV-style)", "Armed cash logistics contract", "Monthly", "Escorted daily banking; audit-registry item"],
    ["Cash Handling", "Till Audit Sheets (Carbonless)", "Daily float in/out, signed", "Monthly", "Reconciles POS vs physical cash"],

    # ============ D. FRONT-OF-HOUSE GUEST SUPPLIES ============
    ["Guest", "Branded Serviettes (2-ply)", "Recycled kraft, 5000ct", "Daily", "Branded wrap; MOQ 100k"],
    ["Guest", "Sugar Sachets (White/Brown/Sweetener)", "Individually wrapped", "Daily", "Condiment bar standard"],
    ["Guest", "Cup Markers (Chalk/Permanent hybrid)", "Branded, 12pk", "Weekly", "Order-cup writing; never run out"],
    ["Guest", "Wooden Stirrers (Branded sleeve)", "FSC birch, 500ct", "Daily", "Hot drink standard"],
    ["Guest", "Straws (Wide-bore, Paper-wrapped)", "Compostable, smoothie-grade", "Daily", "For smoothies (SM-01..04) and lassi"],
    ["Guest", "Wet-Wipes (Branded)", "Lemon scent, alcohol-free", "Daily", "Sticky pancake fingers / bowls"],
    ["Guest", "Napkin Dispensers", "One-at-a-time, branded", "Permanent", "Reduces waste; guest pull standard"],
    ["Guest", "Lids & Cup Sleeves", "Thermal, 8/12/16oz", "Daily", "Hot-drink standard"],
    ["Guest", "Takeaway Carrier Trays (4-cup)", "Moulded fibre, branded", "Daily", "Multi-order carry"],

    # ============ E. FOOD PREP & SERVICE UTENSILS ============
    ["Food Prep", "Pancake Tongs (Silicone-tipped)", "Heat-rated 260°C, long handle", "Permanent", "Handling hot pancakes safely"],
    ["Food Prep", "Latex-Free Gloves (S/M/L)", "Powder-free, textured", "Daily", "Food-contact standard; stock all sizes"],
    ["Food Prep", "Disposable Poly-Aprons", "20-micron clear", "Daily", "Overlay for pancake/bowl prep"],
    ["Food Prep", "Beard Covers & Hairnets", "Breathable, branded band", "Daily", "Food handler requirement"],
    ["Food Prep", "Batter Mixing Buckets (10L/20L)", "NSF HDPE, graduated, lidded", "Permanent", "Pancake batter consistency"],
    ["Food Prep", "Whisks, Spatulas, Bench Scrapers", "Commercial 304 stainless", "Permanent", "BOH toolkit"],
    ["Food Prep", "Digital Thermometers", "Instant-read, NSF", "Permanent", "Milk/food temp logs"],
    ["Food Prep", "Kitchen Timers (Magnetic)", "Rugged digital, loud", "Permanent", "Steep/brew timing"],
    ["Food Prep", "Food-Safe Markers (Date-code)", "Industrial black/red", "Monthly", "Use-By dating every opened consumable"],
    ["Food Prep", "Blue Detectable Plasters", "Food industry standard", "Monthly", "First-aid + contamination protocol"],
    ["Food Prep", "Ice Scoops & Bean Scoops", "Stainless, branded handle", "Permanent", "Portion control"],

    # ============ F. CLEANING & SANITATION (McDonald's bucket system) ============
    ["Sanitation", "Sanitiser Buckets (Red) & Soap Buckets (Green)", "Colour-coded 5L + lids", "Permanent", "Cross-contamination prevention"],
    ["Sanitation", "Sanitiser Solution (Quat tablets)", "SABS approved, bulk", "Weekly", "3-compartment sink standard"],
    ["Sanitation", "Dish Soap (Antibacterial, concentrate)", "Bulk 5L refill", "Weekly", "Smallwares soak + handwash"],
    ["Sanitation", "Degreaser (Food-safe spray + 5L refill)", "SABS approved", "Weekly", "Counter and machine surfaces"],
    ["Sanitation", "Sanitiser Test Strips (QAC/Chlorine)", "100pk", "Monthly", "Health-log mandatory verification"],
    ["Sanitation", "Grease Trap Treatment (Bio-enzymatic)", "Concentrated, monthly dose", "Monthly", "Drainage maintenance — kiosk plumbing"],
    ["Sanitation", "Glass/Window Cleaner", "Streak-free, 5L", "Monthly", "32\" window screen + kiosk fibreglass"],
    ["Sanitation", "Microfibre Cloths (4-colour set)", "Colour-coded by surface", "Bi-monthly", "Red=toilet, green=counter, blue=glass, yellow=general"],
    ["Sanitation", "Heavy-Duty Bin Liners (30L/50L)", "Puncture-resistant, recycled", "Daily", "Kiosk bins"],
    ["Sanitation", "Industrial Mop & Bucket Set", "Colour-coded wringer", "Permanent", "Floor spill + close routine"],
    ["Sanitation", "Hand Sanitiser (High-alcohol)", "Wall-mount + counter", "Monthly", "Staff + guest touchpoints"],
    ["Sanitation", "Jura Cleaning Tablets / Milk-line detergent / Descaler", "OEM verified", "Weekly", "Machine care — per audit registry"],

    # ============ G. INGREDIENTS — kiosk menu (per ACCIO intake) ============
    ["Ingredients", "Fruit Bowl & Smoothie Fruit (mango, pineapple, kiwi, berries, melon, banana, apple, lime, cucumber, mint, ginger, spinach, dragon fruit)", "Fresh, daily wholesale", "Daily", "Gauteng fruit & veg supplier — ACCIO Task B"],
    ["Ingredients", "Yoghurt, Almond Butter, Chia Seeds, Whey Protein", "Bulk tubs", "Weekly", "Smoothie/custom-bowl bases"],
    ["Ingredients", "Pancake Dry Mix (house recipe)", "Pre-mix bags", "Weekly", "Consistency across portals"],
    ["Ingredients", "Eggs, Milk (full/low), Cinnamon, Castor Sugar", "Bulk", "Weekly", "Bush Pancake Stack (R65)"],
    ["Ingredients", "Honey, Granola, Coconut, Mixed Nuts, Cacao Nibs", "Bulk", "Weekly", "Custom bowl toppings"],

    # ============ H. EQUIPMENT & CONSUMABLES (coffee + tech) ============
    ["Equipment", "Espresso Machine (≤40cm depth, semi/super-auto)", "Target < $1,200/unit", "Permanent", "ACCIO Task A — spec + source"],
    ["Equipment", "Burr Grinder (programmable dose)", "Commercial", "Permanent", "ACCIO Task A"],
    ["Equipment", "Water Filtration System", "Countertop, municipal", "Permanent", "ACCIO Task A; Jura cartridges monthly"],
    ["Equipment", "Industrial 5G/LTE Router + Redundant SIMs", "High-thermal rated", "Permanent", "VaultMesh telemetry + POS connectivity"],
    ["Equipment", "1KVA UPS + Surge Protection", "Industrial", "Permanent", "Power blip protection — 24/7"],
    ["Equipment", "CCTV Standalone Unit", "Kiosk-mounted", "Permanent", "Security + incident evidence"],
    ["Equipment", "Biometric Attendance Terminal", "Facial/fingerprint", "Permanent", "Shift clock-in/out — audit registry"],

    # ============ I. STAFF PPE / UNIFORM (2026-2030) ============
    ["Staff", "Polo Shirts (tech-breathable)", "Branded mint/teal", "Annual", "2× per staff member"],
    ["Staff", "Aprons (branded)", "Washable, 2× per staff", "Annual", "Brand layer"],
    ["Staff", "Caps (branded structured)", "3D embroidered logo", "Annual", "1× per staff"],
    ["Staff", "Vellies (Veldskoen, embossed)", "SA leather stitch-down", "Annual", "1× pair per staff"],
    ["Staff", "Steel-Toe Safety Shoes", "SABS approved", "Annual", "Logistics/setup phases"],
    ["Staff", "High-Visibility Vests", "2× per kiosk", "Annual", "Setup + night service"],
    ["Staff", "Digital e-Ink Nametags (NFC)", "base44 IAM linked", "Permanent", "ID + kiosk access key"],
    ["Staff", "First-Aid Kit (SABS/H&S) + refills", "Wall-mount", "Monthly", "Health compliance"],

    # ============ J. ADMIN / CLOSING / COMPLIANCE ============
    ["Admin", "Operational Logbooks + Digital Writing Pads", "e-Ink sync to ComplianceAudit", "Quarterly", "Hygiene, cleaning, shift handover logs"],
    ["Admin", "Closing Checklists (laminated, SOP v1.2)", "Dry-erase + markers", "Permanent", "Close routine 365 days"],
    ["Admin", "Receipt Thermal Rolls (spare case)", "80mm, BPA-free", "Quarterly", "Never run out — bank records"],
    ["Admin", "Archive Folders (7-year retention)", "Lever-arch, branded", "Annual", "Tax/legal records per registry"],
    ["Admin", "Petty Cash Vouchers + Pencil Case kit", "Finance pack", "Quarterly", "Small cash ops"],
    ["Admin", "Lock-out/Tag-out (LOTO) Kit", "OSHA standard", "Permanent", "Machine maintenance safety"],
    ["Admin", "Municipal Trading Permit + H&S Register + COIDA + Fire Cert + Food Handler medicals", "Annual renewal", "Annual", "Legal checklist — per audit registry"],

    # ============ K. LOGISTICS & SUPPLY ============
    ["Logistics", "Loading Trolley (Hand Truck)", "Folding, 150kg", "Permanent", "Daily stock delivery"],
    ["Logistics", "Shipping Pallets & Labels", "Bulk", "Monthly", "Stock intake"],
    ["Logistics", "Fleet/Fuel Cards + Mileage Logbooks", "Vehicle ops", "Quarterly", "Restock runs"],
    ["Logistics", "High-Security Padlocks + RFID Access Tags", "Kiosk + storage", "Annual", "Asset security"],
    ["Logistics", "Toll-gate (e-Tag) Transponders", "Vehicle", "As-needed", "Restock runs"],
]

cols = ["Category", "Item", "Specification / Standard", "Order Freq", "Notes / Agent Task"]

wb = Workbook()
ws = wb.active; ws.title = 'Fruitful BOQ - 365x24x7'
ws.sheet_properties.tabColor = HEADER_FILL

thin = Side(style='thin', color=BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name=FONT, size=11, bold=True, color=HEADER_TXT)
header_fill = PatternFill('solid', start_color=HEADER_FILL)
band_fill = PatternFill('solid', start_color=BAND)

for j, name in enumerate(cols, 1):
    c = ws.cell(1, j, name); c.font = header_font; c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border

for i, row_data in enumerate(data, start=2):
    for j, val in enumerate(row_data, 1):
        c = ws.cell(i, j, val); c.border = border
        if (i % 2) == 1: c.fill = band_fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

widths = [18, 45, 32, 16, 45]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws.freeze_panes = 'A2'
wb.save('Fruitful_Portal_BOQ_Complete_365x24x7.xlsx')
print(f"DONE: {len(data)} line items")
