# master_boq_v1.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- theme ---
FONT, HEADER_FILL, HEADER_TXT, BAND, BORDER = 'Arial', '0FB5A6', 'FFFFFF', 'F2F2F2', 'BFBFBF'

data = [
    # --- SMALLWARES & TOOLS (The Starbucks/McDonald's standard) ---
    ["Smallwares", "Milk Frothing Pitchers (350ml/600ml)", "Branded Stainless Steel", "Daily", "Source 304 Stainless with Fruitful logo engraving"],
    ["Smallwares", "Coffee Tampers & Mats", "Weighted Pro-Series", "Permanent", "Source 58mm calibrated tampers"],
    ["Smallwares", "Digital Thermometers", "NSF Certified", "Daily", "Instant-read probe thermometers"],
    ["Smallwares", "Pancake Batter Dispensers", "Commercial Trigger-action", "Daily", "Stainless steel gravity-fed units"],
    ["Smallwares", "Silicone Spatulas & Spreaders", "High-Heat Rated", "Daily", "Heat resistant up to 260°C"],
    ["Smallwares", "Scoops & Tongs Set", "Branded Handle", "Daily", "Ice scoops, bean scoops, food tongs"],

    # --- BRANDED DISPOSABLES (The volume movers) ---
    ["Disposables", "Branded Serviettes (Napkins)", "Recycled Kraft 2-ply", "Daily", "MOQ 100k; 4-color print"],
    ["Disposables", "Branded Cup Markers", "Chalk/Permanent Hybrid", "Weekly", "Specialized markers for writing on paper/plastic cups"],
    ["Disposables", "Cup Carriers (4-cup)", "Moulded Fiber", "Daily", "Biodegradable, Fruitful embossed"],
    ["Disposables", "Wooden Stirrers", "Branded Wrap", "Daily", "FSC certified birch wood"],

    # --- CLEANING & FACILITY (Deep maintenance) ---
    ["Cleaning", "Sanitizer Test Strips", "QAC/Chlorine", "Daily", "Mandatory for health compliance logs"],
    ["Cleaning", "Grease Trap Treatment", "Bio-Enzymatic", "Monthly", "Concentrated liquid for drainage maintenance"],
    ["Cleaning", "Milk System Cleaners", "Jura / OEM Equivalent", "Daily", "Specific for automatic milk frothers"],
    ["Cleaning", "Heavy-Duty Buckets & Mops", "Industrial Grade", "Permanent", "Color-coded for cross-contamination prevention"],
    ["Facility", "Branded Entrance Mats", "Anti-fatigue / Logo", "Permanent", "High-traffic water-absorbent mats"],
    ["Facility", "Kiosk Barrier Poles", "Teal Powder-coated", "Permanent", "Retractable belt stanchions for queue management"],

    # --- STAFF TECH & OPERATIONS (HotStack / base44) ---
    ["Staff Tech", "Digital Nametags (e-Ink)", "API-enabled", "Permanent", "4-color e-Ink with NFC access key"],
    ["Staff Tech", "Rugged 10\" Tablets", "base44 SDK Pre-loaded", "Permanent", "Industrial casing with 5G connectivity"],
    ["Staff Tech", "Digital Writing Pads", "e-Ink Logbooks", "Permanent", "For digitizing hygiene and H&S logs"],
    ["Staff Tech", "5G Industrial Routers", "VaultMesh Node", "Permanent", "High-thermal rated for kiosk enclosures"],

    # --- CONSUMABLES (The high-quality core) ---
    ["Consumables", "Artisan Coffee Beans", "GreenCoffeeReceipt Grade", "Daily", "Fruitful Certified (SCA 80+ / Moisture 9-12%)"],
    ["Consumables", "UHT Milk (Full/Low Fat)", "Local Supply", "Daily", "Bulk procurement agreement"],
    ["Consumables", "Artisan Syrups", "Glass Bottles / Branded", "Weekly", "Vanilla, Caramel, Hazelnut"],
    ["Consumables", "Pancake Mix", "Fruitful Signature", "Weekly", "Pre-mix dry bags for consistency"],

    # --- APPAREL (The 2026-2030 Line) ---
    ["Apparel", "Polo Shirts", "Tech-Breathable", "Annual", "Branded Mint-Green/Teal"],
    ["Apparel", "Bomber Jackets", "VaultMesh Weatherproof", "Annual", "Insulated for early/late shifts"],
    ["Apparel", "Vellies (Veldskoen)", "Embossed Leather", "Permanent", "Traditional SA boots with brand mark"],
    ["Apparel", "Digital Nametags", "NFC / IAM Linked", "Permanent", "Dynamic display tags"]
]

cols = ["Category", "Item Description", "Specification / Standard", "Order Freq", "Agent Sourcing Task"]

wb = Workbook()
ws = wb.active; ws.title = 'Fruitful Master BOQ'
ws.sheet_properties.tabColor = HEADER_FILL

thin = Side(style='thin', color=BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name=FONT, size=11, bold=True, color=HEADER_TXT)
header_fill = PatternFill('solid', start_color=HEADER_FILL)
band_fill   = PatternFill('solid', start_color=BAND)

for j, name in enumerate(cols, 1):
    c = ws.cell(1, j, name); c.font = header_font; c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border

for i, row_data in enumerate(data, start=2):
    for j, val in enumerate(row_data, 1):
        c = ws.cell(i, j, val); c.border = border
        if (i % 2) == 1: c.fill = band_fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

for j, name in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(j)].width = 35

ws.freeze_panes = 'A2'
wb.save('Fruitful_Portal_Master_BOQ_Enhanced.xlsx')
print("DONE")
