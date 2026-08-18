# gen_menu_boq_engine.py — Radical menu cost engine: every grain, McDonald's + Mercedes + scale doctrine
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont = Font(name=FONT, size=11, bold=True, color='FFFFFF')
hfill = PatternFill('solid', start_color='0FB5A6')
bfill = PatternFill('solid', start_color='F2F2F2')
gold = PatternFill('solid', start_color='FFF2CC')
green = PatternFill('solid', start_color='E2F0D9')

def sheet(ws, headers, widths):
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h); c.font = hfont; c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'

def put(ws, r, vals, fill=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = border
        if r % 2 == 1 and fill is None: c.fill = bfill
        if fill: c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical='top')

def money(v): return f"R{v:,.2f}"

wb = Workbook()

# ============ Sheet 1: The 8 lines, every grain ============
ws = wb.active; ws.title = '8 Lines - Every Grain'
ws.sheet_properties.tabColor = '0FB5A6'
sheet(ws, ["Line / component", "Qty/unit", "Unit cost", "R per unit", "Flag"], [40, 12, 14, 12, 30])
rows = [
    # coffee with milk
    ("☕ THE COFFEE WITH MILK (R38) — 350ml standard", "", "", "", ""),
    ("  Espresso beans (18g @ R348/kg)", "18 g", "R0.348/g", 6.26, "Factual (Aquazania R1,392/4kg)"),
    ("  Milk (150ml @ R22.04/L)", "150 ml", "R22.04/L", 3.31, "Factual (Dewfresh R132.25/6L shop)"),
    ("  Cup 12oz ripple branded", "1", "R2.50", 2.50, "Estimate (wholesale)"),
    ("  Lid", "1", "R1.00", 1.00, "Estimate"),
    ("  Sleeve", "1", "R0.80", 0.80, "Estimate"),
    ("  Stirrer", "1", "R0.15", 0.15, "Estimate"),
    ("  Serviette", "1", "R0.30", 0.30, "Estimate"),
    ("  Sugar sachet", "1", "R0.34", 0.34, "Factual (Huletts R338.10/1000)"),
    ("  COFFEE DIRECT COST", "", "", 14.66, ""),
    ("  Gross margin @ R38", "", "", 23.34, "61.4%"),
    # froth = milk + foam (same stack, +0)
    ("☕ THE COFFEE WITH FROTH (R38) — same stack", "", "", 14.66, "foam = milk, no extra cost"),
    ("☕ THE COFFEE WITH CHOCOLATE (R38) — +8g cocoa", "8 g", "R0.30/g", 2.40, "Cocoa powder Dutch 1kg ~R300 (est.)"),
    ("  Chocolate line direct cost", "", "", 17.06, ""),
    ("  Gross margin @ R38", "", "", 20.94, "55.1%"),
    # smoothies
    ("🥤 SMOOTHIE (R55) — 400ml one size", "", "", "", ""),
    ("  Ingredients (blended avg: green/lassi/protein)", "400 ml", "—", 18.00, "Estimate (protein line higher, lassi lower)"),
    ("  Cup 400ml clear PET", "1", "R3.00", 3.00, "Estimate"),
    ("  Dome lid", "1", "R2.00", 2.00, "Estimate"),
    ("  Straw wide", "1", "R0.40", 0.40, "Estimate"),
    ("  SMOOTHIE DIRECT COST", "", "", 23.40, ""),
    ("  Gross margin @ R55", "", "", 31.60, "57.5%"),
    # pancakes
    ("🥞 PANNEKOEK (traditional SA, R35 x n) — 1 piece boxed", "", "", "", ""),
    ("  Cake flour 23g (250g/11 crepes)", "23 g", "R0.02/g", 0.45, "Traditional 2-2-2 recipe"),
    ("  Egg (0.18)", "0.18", "R2.40/egg", 0.44, "2 eggs per 11-crepe batch"),
    ("  Water 45ml", "45 ml", "tap", 0.00, "Traditional recipe: no milk"),
    ("  Baking powder + salt", "0.9 g", "—", 0.03, "Estimate"),
    ("  Pan oil (2ml)", "2 ml", "R40/L", 0.08, "Estimate"),
    ("  Cinnamon sugar (8g + 1.5g)", "9.5 g", "—", 0.16, "Estimate"),
    ("  Lemon wedge", "1/8", "R2.50/lemon", 0.31, "Estimate"),
    ("  Box (prepacked)", "1", "R1.25", 1.25, "Local wholesale (GZ R0.28)"),
    ("  PANNEKOEK DIRECT COST / UNIT", "", "", 2.72, ""),
    ("  Gross margin @ R35", "", "", 32.28, "92.2%"),
    ("  Pancake 12-pack (R420)", "12", "R2.72", 32.64, "margin R387.36 (92.2%)"),
]
for i, r in enumerate(rows, 2):
    if r[0].startswith(("☕", "🥤", "🥞")) and not r[0].startswith("  "):
        put(ws, i, list(r), fill=green)
    else:
        put(ws, i, list(r), fill=gold if r[0].strip().endswith("COST") or "margin" in r[0].lower() else None)

# ============ Sheet 2: Monthly menu P&L (per portal, neighbourhood) ============
ws2 = wb.create_sheet('Monthly Menu P&L')
sheet(ws2, ["Line", "Coffee (1,500/mo)", "Smoothies (750/mo)", "Pancakes (600 units/mo)", "TOTAL"], [30, 16, 16, 16, 14])
COF_U = 14.66; SM_U = 23.40; PAN_U = 2.72  # traditional pannekoek
COF_Q, SM_Q, PAN_Q = 1500, 750, 600
rows2 = [
    ("Units", COF_Q, SM_Q, PAN_Q, COF_Q + SM_Q + PAN_Q),
    ("Revenue", COF_Q*38, SM_Q*55, PAN_Q*35, COF_Q*38 + SM_Q*55 + PAN_Q*35),
    ("Direct COGS (grains above)", COF_Q*COF_U, SM_Q*SM_U, PAN_Q*PAN_U, COF_Q*COF_U + SM_Q*SM_U + PAN_Q*PAN_U),
    ("COGS %", COF_U/38, SM_U/55, PAN_U/35, (COF_Q*COF_U + SM_Q*SM_U + PAN_Q*PAN_U)/(COF_Q*38 + SM_Q*55 + PAN_Q*35)),
    ("Payment fees 3%", COF_Q*38*0.03, SM_Q*55*0.03, PAN_Q*35*0.03, (COF_Q*38 + SM_Q*55 + PAN_Q*35)*0.03),
    ("Labour (2 staff, mgr=HQ)", 8842.00, 4421.00, 3537.00, 16800.00),
    ("Mall rent + Jura + opex", "", "", "", 12082.00),
    ("Royalty 6% + marketing 4%", "", "", "", (COF_Q*38 + SM_Q*55 + PAN_Q*35)*0.10),
    ("EBITDA", "", "", "", None),
    ("Depreciation + tax (27%)", "", "", "", None),
    ("NET PROFIT / month", "", "", "", None),
]
for i, r in enumerate(rows2, 2):
    if r[0] == "COGS %":
        put(ws2, i, [r[0]] + [f"{v*100:.1f}%" for v in r[1:4]] + [f"{r[4]*100:.1f}%"])
    elif all(isinstance(v, (int, float)) for v in r[1:]):
        put(ws2, i, [r[0]] + [money(v) for v in r[1:4]] + [money(r[4])])
    else:
        put(ws2, i, list(r))
rev = COF_Q*38 + SM_Q*55 + PAN_Q*35
cogs = COF_Q*COF_U + SM_Q*SM_U + PAN_Q*PAN_U
ebitda = rev - cogs - rev*0.03 - 16800 - 12082 - rev*0.10
net = (ebitda - 4167) * 0.73
put(ws2, 10, ["EBITDA", "", "", "", money(ebitda)], fill=gold)
put(ws2, 11, ["Depreciation + tax (27%)", "", "", "", f"-R{4167 + (ebitda-4167)*0.27:,.0f}"])
put(ws2, 12, ["NET PROFIT / month", "", "", "", money(net)], fill=gold)
hn = ws2.cell(13, 1, "Honest note: radical menu at baseline volume = thinner revenue (R119,250 vs R198,260 full menu) but zero waste, hours-long training, and the BushCoding funnel. Levers: volume (jam study), 1.5-staff model, pancake 12-packs, BushCoding revenue.")
hn.font = Font(italic=True, color='666666')

# ============ Sheet 3: McDonald's / Mercedes / Jack Ma doctrine ============
ws3 = wb.create_sheet('Doctrine')
sheet(ws3, ["Layer", "Rule", "Applied to Fruitful"], [26, 46, 40])
doc = [
    ("McDonald's", "Cut to the 90% of sales; quarterly re-cut by data", "8 lines only; VaultMesh trim report cans bottom-10% every quarter; Pitaya Punch is the designated replacement"),
    ("McDonald's", "One spec, one recipe, one till — training in hours", "8 buttons, 3 prices, 1 stepper; no modifiers anywhere"),
    ("McDonald's", "Prepack = zero close-of-day waste", "Pancake unit boxed; 12-pack is the DC unit"),
    ("Mercedes", "One part number per SKU, identical across the fleet", "MB-style P/N: FFP-01..08 (see Sheet 4); same part in all 300 portals"),
    ("Mercedes", "No options, no trim levels, no half measures", "R450k package = one spec; menu = one spec; training = one course"),
    ("Jack Ma", "Scale the system, not the effort", "Amazon DC supply spine; NUCO fleet for deploy only; quarterly automation"),
    ("Jack Ma", "Start with the customer's 20 seconds, compound the rest", "order-to-pay <30s; freed attention -> BushCoding sessions"),
    ("Jack Ma", "If it can't be measured, it doesn't exist", "per-pour ml telemetry, order times, trim reports, waste = 0 target"),
]
for i, d in enumerate(doc, 2): put(ws3, i, list(d))

# ============ Sheet 4: Part numbers (Mercedes) ============
ws4 = wb.create_sheet('Part Numbers')
sheet(ws4, ["P/N", "SKU", "Spec (one version only)", "Menu price", "Direct cost", "Margin %", "MOQ / contract"], [10, 30, 36, 12, 12, 10, 26])
pn = [
    ("FFP-01", "The Coffee with Milk", "350ml, beans 18g, milk 150ml", "R38", 14.66, "61.4%", "beans: 27kg/mo @ R348/kg (Aquazania)"),
    ("FFP-02", "The Coffee with Froth", "350ml, +foam", "R38", 14.66, "61.4%", "same bean contract"),
    ("FFP-03", "The Coffee with Chocolate", "350ml, +8g cocoa", "R38", 17.06, "55.1%", "cocoa: 1kg/mo (est. R300)"),
    ("FFP-04", "Green Energy smoothie", "400ml one size", "R55", 23.40, "57.5%", "ingredients kit weekly (est. R2,500)"),
    ("FFP-05", "Mango Lassi smoothie", "400ml one size", "R55", 23.40, "57.5%", "weekly kit"),
    ("FFP-06", "Power Protein smoothie", "400ml one size", "R55", 23.40, "57.5%", "weekly kit + whey monthly"),
    ("FFP-07", "Pannekoek unit (prepacked)", "1pc boxed, cinnamon sugar + lemon (traditional)", "R35", 2.72, "92.2%", "packer contract; 12-pack = PANC-12"),
    ("FFP-08", "Pannekoek 12-pack (DC unit)", "12 units, sealed box", "R420", 32.64, "92.2%", "Amazon-DC listing PANC-12"),
]
for i, p in enumerate(pn, 2): put(ws4, i, list(p))

wb.save('FRUITFUL_MENU_BOQ_ENGINE.xlsx')
print("DONE")
